"""Module 2 - text extraction from PDF / image / Word / Excel / text files."""
from __future__ import annotations

from pathlib import Path

ALLOWED_EXTENSIONS = {".pdf", ".png", ".jpg", ".jpeg", ".docx", ".xlsx", ".txt"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
MAX_FILE_MB = 10


def validate_file(filename: str, size_bytes: int) -> None:
    """Module 1 validation. Raises ValueError with a friendly message."""
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{ext}'. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
        )
    if size_bytes == 0:
        raise ValueError("The file is empty.")
    if size_bytes > MAX_FILE_MB * 1024 * 1024:
        raise ValueError(f"File is larger than {MAX_FILE_MB} MB.")


def ocr_available() -> bool:
    try:
        import pytesseract  # noqa: F401
        from PIL import Image  # noqa: F401
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("This PDF is password protected.")
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n".join(pages)


def _extract_image(path: Path) -> str:
    if not ocr_available():
        return ""
    import pytesseract
    from PIL import Image

    return pytesseract.image_to_string(Image.open(path))


def _extract_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError:
        raise ValueError("Reading Word files needs python-docx (pip install python-docx).")
    d = docx.Document(str(path))
    lines = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            lines.append("  ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def _extract_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError("Reading Excel files needs openpyxl (pip install openpyxl).")
    wb = load_workbook(str(path), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [_cell_text(v) for v in row if v is not None and str(v).strip()]
            if cells:
                # "Label: value" rows read naturally for the regex parser
                lines.append(": ".join(cells) if len(cells) == 2 else "  ".join(cells))
    wb.close()
    return "\n".join(lines)


def _cell_text(v) -> str:
    from datetime import date, datetime
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def extract_text(path: str | Path) -> str:
    """Return the text of a file ('' if nothing could be read)."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        text = _extract_pdf(path)
        if not text.strip() and ocr_available():
            # Scanned PDF: OCR is optional and needs pdf2image + poppler.
            try:
                from pdf2image import convert_from_path
                import pytesseract

                text = "\n".join(pytesseract.image_to_string(img)
                                 for img in convert_from_path(str(path)))
            except Exception:
                pass
        return text
    if ext in IMAGE_EXTENSIONS:
        return _extract_image(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext == ".xlsx":
        return _extract_xlsx(path)
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type '{ext}'")
