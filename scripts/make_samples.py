"""Generate 24 demo documents whose dates are relative to a reference day.

Every reminder case is covered: overdue, due today, critical (<= 7 days),
expiring soon (8-30 days), active (> 30 days), no expiry, a reminder already
marked completed, and an incomplete document with a low health score.
Formats: PDF, Word (.docx), Excel (.xlsx) and plain text.

    python scripts/make_samples.py                  # relative to today
    python scripts/make_samples.py --today 2026-10-15
"""
from __future__ import annotations

import argparse
from datetime import date, timedelta
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "sample_docs"
FOOTER = "SAMPLE DOCUMENT - generated for the LifeAdmin demo"
SAMPLE_EXTS = {".pdf", ".docx", ".xlsx", ".txt"}

# filename -> (expected type, expected status level) for tests and the demo loader
EXPECTED = {}
# filenames whose reminder the demo marks as "completed" (renewed already)
COMPLETED = {"old_passport.pdf"}
# a short note shown on the Documents page
NOTES = {
    "old_passport.pdf": "Renewed - new passport uploaded separately.",
    "car_insurance.pdf": "Compare quotes before renewing.",
    "electricity_bill.pdf": "Pay via TSSPDCL app.",
    "unknown_scan.txt": "Scan was unclear - fill in the missing details.",
}


def d_(today: date, n: int) -> date:
    return today + timedelta(days=n)


def dmy(d: date) -> str:          # 15/10/2026
    return d.strftime("%d/%m/%Y")


def samples(today: date):
    """List of (filename, lines). Lines[0] is the heading."""
    t = lambda n: d_(today, n)
    S = []

    def add(name, typ, level, lines):
        EXPECTED[name] = (typ, level)
        S.append((name, lines))

    # ---------------- OVERDUE (expired)
    add("old_puc.pdf", "Vehicle", "expired", [
        "Pollution Under Control Certificate",
        "Issued By: Government Authorised PUC Centre, Kukatpally",
        "Certificate No: PUC00998877",
        "Vehicle Registration No: TS 09 EA 4521",
        f"Date of Issue: {dmy(t(-200))}",
        f"Valid Until: {dmy(t(-20))}",
    ])
    add("bike_insurance.pdf", "Insurance", "expired", [
        "Bajaj Allianz General Insurance Company Limited",
        "TWO WHEELER PACKAGE POLICY",
        "Policy No: TW4471029",
        "Insured Name: Adityavardhan Reddy",
        "Vehicle Registration No: TS 08 GK 7310",
        f"Period of Insurance: {dmy(t(-375))} to {dmy(t(-10))}",
        "Total Premium: Rs. 1,980.00",
    ])
    add("old_passport.pdf", "ID", "expired", [
        "Republic of India - Passport",
        "Issuing Authority: Passport Office, Hyderabad",
        "Passport No: K4471023",
        "Name: Srinivas Reddy",
        f"Date of Issue: {dmy(t(-3710))}",
        f"Date of Expiry: {dmy(t(-60))}",
    ])
    # ---------------- DUE TODAY
    add("driving_licence.pdf", "ID", "critical", [
        "Transport Department, Government of Telangana - Driving Licence",
        "Issued By: Regional Transport Office, Hyderabad",
        "Licence No: TS0920190012345",
        "Name: Adityavardhan Reddy",
        "Date of Birth: 14/08/2004",
        f"Date of Issue: {dmy(t(-1800))}",
        f"Valid Until: {dmy(t(0))}",
    ])
    # ---------------- CRITICAL (1-7 days)
    add("electricity_bill.pdf", "Bill", "critical", [
        "Southern Power Distribution Company of Telangana Limited",
        "ELECTRICITY BILL",
        "Consumer No: 1102-455-8821",
        "Customer Name: Srinivas Reddy",
        f"Bill Date: {dmy(t(-12))}",
        f"Due Date: {dmy(t(3))}",
        "Units Consumed: 214",
        "Amount Due: Rs. 1,842.00",
    ])
    add("credit_card_bill.pdf", "Bill", "critical", [
        "HDFC Bank Credit Card Statement",
        "Card No: XXXX XXXX XXXX 4417",
        "Customer Name: Adityavardhan Reddy",
        f"Statement Date: {t(-15).strftime('%d %b %Y')}",
        f"Payment Due Date: {t(5).strftime('%d %b %Y')}",
        "Total Amount Due: Rs. 12,460.00",
    ])
    add("netflix_subscription.xlsx", "Subscription", "critical", [
        "Netflix Subscription",
        ("Membership No", "NFX-99812"),
        ("Customer Name", "Adityavardhan Reddy"),
        ("Plan", "Standard - monthly plan"),
        ("Next Billing Date", t(6)),
        ("Amount", "Rs. 499"),
    ])
    # ---------------- EXPIRING SOON (8-30 days)
    add("car_insurance.pdf", "Insurance", "warning", [
        "ICICI Lombard General Insurance Company Limited",
        "PRIVATE CAR PACKAGE POLICY SCHEDULE",
        "Policy No: INS784521",
        "Insured Name: Srinivas Reddy",
        "Vehicle Registration No: TS 09 EA 4521",
        f"Period of Insurance: {dmy(t(-353))} to {dmy(t(12))}",
        "Total Premium: Rs. 14,250.00",
    ])
    add("broadband_bill.docx", "Bill", "warning", [
        "ACT Fibernet - Broadband Bill",
        "Account No: ACT20918377",
        "Customer Name: Srinivas Reddy",
        f"Bill Date: {t(-5).strftime('%d-%m-%Y')}",
        f"Due Date: {t(15).strftime('%d-%m-%Y')}",
        "Plan: 150 Mbps monthly plan",
        "Total Payable: Rs. 1,179.00",
    ])
    add("gym_membership.txt", "Subscription", "warning", [
        "Cult.fit Gym Membership",
        "Membership No: CULT-55120",
        "Name: Adityavardhan Reddy",
        "Plan: Quarterly membership",
        f"Renewal Date: {t(22).strftime('%B %d, %Y')}",
        "Fee: Rs. 6,999",
    ])
    add("vehicle_fitness.pdf", "Vehicle", "warning", [
        "Transport Department - Fitness Certificate for Motor Vehicle",
        "Issued By: Regional Transport Authority, Hyderabad",
        "Certificate No: FC2025-77812",
        "Vehicle Registration No: TS 09 UB 2290",
        f"Date of Issue: {dmy(t(-337))}",
        f"Valid Until: {dmy(t(28))}",
    ])
    # ---------------- ACTIVE (> 30 days)
    add("phone_warranty.pdf", "Warranty", "active", [
        "Samsung India Electronics - Warranty Card",
        "Product: Galaxy S24, Model No: SM-S921B",
        "Serial Number: R5CX21ABCDE",
        "Customer Name: Adityavardhan Reddy",
        f"Date of Purchase: {dmy(t(-317))}",
        f"Warranty valid until {dmy(t(48))}",
        "Price: Rs. 74,999",
    ])
    add("domain_hosting.pdf", "Subscription", "active", [
        "GoDaddy - Domain Renewal Notice",
        "Order No: GD-88123004",
        "Customer Name: Adityavardhan Reddy",
        "Domain: redorca.in - annual plan, auto-renew off",
        f"Renewal Date: {t(75).strftime('%Y-%m-%d')}",
        "Renewal Price: Rs. 899",
    ])
    add("property_tax.pdf", "Bill", "active", [
        "Greater Hyderabad Municipal Corporation - Property Tax Bill",
        "Bill No: PTIN-1180044712",
        "Owner Name: Srinivas Reddy",
        f"Bill Date: {dmy(t(-20))}",
        f"Due Date: {dmy(t(95))}",
        "Amount Payable: Rs. 8,420.00",
    ])
    add("health_insurance.pdf", "Insurance", "active", [
        "Star Health and Allied Insurance Co. Ltd",
        "FAMILY HEALTH OPTIMA - POLICY CERTIFICATE",
        "Policy Number: HLT556201",
        "Policy Holder: Lakshmi Reddy",
        "Sum Insured: Rs. 5,00,000",
        f"Start Date: {dmy(t(-200))}",
        f"Expiry Date: {dmy(t(165))}",
        "Premium: Rs. 21,800",
    ])
    add("life_insurance_lic.docx", "Insurance", "active", [
        "Life Insurance Corporation of India",
        "Policy No: LIC4455667",
        "Policy Holder: Srinivas Reddy",
        "Plan: Jeevan Anand",
        f"Renewal Date: {t(240).strftime('%d %B %Y')}",
        "Premium: Rs. 12,400",
    ])
    add("laptop_warranty.pdf", "Warranty", "active", [
        "Dell Technologies Warranty Card",
        "Product Warranty Certificate",
        "Serial Number: DLX9921877",
        "Customer Name: Adityavardhan Reddy",
        f"Date of Purchase: {dmy(t(-400))}",
        f"Warranty valid until {dmy(t(330))}",
        "Price: Rs. 62,990",
    ])
    add("ac_warranty.xlsx", "Warranty", "active", [
        "Voltas Air Conditioner - Warranty Card",
        ("Serial Number", "VLT15T883201"),
        ("Customer Name", "Lakshmi Reddy"),
        ("Date of Purchase", t(-130)),
        ("Warranty valid until", t(600)),
        ("Price", "Rs. 38,500"),
    ])
    add("vehicle_rc.pdf", "Vehicle", "active", [
        "Registration Certificate (RC) - Motor Vehicle",
        "Issuing Authority: Regional Transport Authority, Hyderabad",
        "Registration No: TS 09 EA 4521",
        "Registered Owner: Srinivas Reddy",
        "Document No: RC2019TS44120",
        "Chassis No: MA3EYD32S00123456   Engine No: K12MN4123456",
        f"Date of Registration: {dmy(t(-2500))}",
        f"Valid Upto: {dmy(t(3000))}",
    ])
    add("passport.pdf", "ID", "active", [
        "Republic of India - Passport",
        "Issuing Authority: Passport Office, Hyderabad",
        "Passport No: W8812034",
        "Name: Adityavardhan Reddy",
        f"Date of Issue: {dmy(t(-50))}",
        f"Date of Expiry: {dmy(t(3600))}",
    ])
    # ---------------- NO EXPIRY
    add("aadhaar_card.pdf", "ID", "none", [
        "Unique Identification Authority of India - Aadhaar",
        "Government of India",
        "Name: Adityavardhan Reddy",
        "Date of Birth: 14/08/2004",
        "Aadhaar No: 2345 6789 0123",
    ])
    add("pan_card.pdf", "ID", "none", [
        "Income Tax Department - PAN Card",
        "Permanent Account Number: ABCPR1234F",
        "Name: Adityavardhan Reddy",
        "Date of Birth: 14/08/2004",
    ])
    add("degree_certificate.pdf", "Certificate", "none", [
        "Osmania University",
        "DEGREE CERTIFICATE",
        "This is to certify that",
        "Name of Candidate: Srinivas Reddy",
        "has been awarded the degree of Bachelor of Commerce",
        "Certificate No: OU1998BC00412",
        f"Date of Issue: {dmy(t(-9500))}",
    ])
    # ---------------- INCOMPLETE (low health score, needs review)
    add("unknown_scan.txt", "Other", "warning", [
        "scanned page 1",
        "valid till " + dmy(t(18)),
        "ref 77812",
    ])
    return S


# ---------------------------------------------------------------- writers
def _text_lines(lines):
    out = []
    for ln in lines:
        if isinstance(ln, tuple):
            k, v = ln
            out.append(f"{k}: {dmy(v) if isinstance(v, date) else v}")
        else:
            out.append(ln)
    return out


def write_pdf(path: Path, lines) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    y = 800
    for i, line in enumerate(_text_lines(lines)):
        c.setFont("Helvetica-Bold" if i == 0 else "Helvetica", 15 if i == 0 else 11)
        c.drawString(60, y, line)
        y -= 30 if i == 0 else 20
    c.setFont("Helvetica-Oblique", 8)
    c.drawString(60, 60, FOOTER)
    c.save()


def write_docx(path: Path, lines) -> None:
    import docx

    d = docx.Document()
    text = _text_lines(lines)
    d.add_heading(text[0], level=1)
    for line in text[1:]:
        d.add_paragraph(line)
    d.add_paragraph(FOOTER)
    d.save(str(path))


def write_xlsx(path: Path, lines) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Document"
    for ln in lines:
        ws.append(list(ln) if isinstance(ln, tuple) else [ln])
    ws.append([FOOTER])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 30
    wb.save(str(path))


def write_txt(path: Path, lines) -> None:
    path.write_text("\n".join(_text_lines(lines) + ["", FOOTER]) + "\n", encoding="utf-8")


WRITERS = {".pdf": write_pdf, ".docx": write_docx, ".xlsx": write_xlsx, ".txt": write_txt}


def main(today: date, out: Path = OUT) -> list:
    """Write every sample into `out` and return their paths."""
    out.mkdir(parents=True, exist_ok=True)
    paths = []
    for name, lines in samples(today):
        p = out / name
        WRITERS[p.suffix](p, lines)
        paths.append(p)
    return paths


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--today", type=date.fromisoformat, default=date.today())
    ref = ap.parse_args().today
    written = main(ref)
    print(f"Wrote {len(written)} sample documents to {OUT} (reference day {ref})")
